"""pdf-to-pptx 测试。

不依赖任何真实客户材料 —— 用 tests/fixtures/sample-deck.pdf（合成的虚构方案）。
断言的是「修补机制生效了」，不是「像素级还原」，所以在任何字体环境下都稳定。
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
CONVERT = ROOT / "scripts" / "convert.py"
FIXTURE = ROOT / "tests" / "fixtures" / "sample-deck.pdf"
FIXTURE_PAGES = 5

sys.path.insert(0, str(ROOT / "scripts"))
import convert as C  # noqa: E402


# ------------------------------------------------------------ 纯函数单测
@pytest.mark.parametrize(
    "raw,expected",
    [
        ("MicrosoftYaHei", "Microsoft YaHei"),
        ("ABCDEF+MicrosoftYaHei", "Microsoft YaHei"),   # 剥子集前缀
        ("TimesNewRoman", "Times New Roman"),
        ("PingFangSC", "PingFang SC"),
        ("SimSun", "SimSun"),        # 真名本来就没空格 —— 不许动
        ("SimHei", "SimHei"),
        ("KaiTi", "KaiTi"),
        ("FangSong", "FangSong"),
        ("Arial", "Arial"),
        (None, None),
    ],
)
def test_norm_font(raw, expected):
    assert C.norm_font(raw) == expected


@pytest.mark.parametrize(
    "text,is_cjk",
    [("中文", True), ("hello", False), ("mixed 中文", True), ("", False), ("１２３", False)],
)
def test_cjk_detection(text, is_cjk):
    assert bool(C.CJK.search(text)) is is_cjk


# ------------------------------------------------- 端到端（需要 LibreOffice）
needs_soffice = pytest.mark.skipif(
    not any(
        (Path(c).exists() if c.startswith("/") else __import__("shutil").which(c))
        for c in C.SOFFICE_CANDIDATES
    ),
    reason="需要 LibreOffice",
)


@pytest.fixture(scope="module")
def converted(tmp_path_factory):
    out = tmp_path_factory.mktemp("out") / "deck.pptx"
    subprocess.run([sys.executable, str(CONVERT), str(FIXTURE), "-o", str(out)],
                   check=True, capture_output=True, timeout=600)
    return Presentation(str(out))


@pytest.fixture(scope="module")
def raw_converted(tmp_path_factory):
    """未修补的对照组，用来证明修补真的改变了什么。"""
    out = tmp_path_factory.mktemp("raw") / "deck.pptx"
    subprocess.run([sys.executable, str(CONVERT), str(FIXTURE), "-o", str(out), "--no-postfix"],
                   check=True, capture_output=True, timeout=600)
    return Presentation(str(out))


@needs_soffice
def test_page_count_preserved(converted):
    assert len(converted.slides._sldIdLst) == FIXTURE_PAGES


@needs_soffice
def test_text_is_editable_not_an_image(converted):
    """核心承诺：文字是可编辑文本框，不是整页截图。"""
    boxes = chars = pictures = 0
    for slide in converted.slides:
        for sh in slide.shapes:
            if "PICTURE" in str(sh.shape_type or ""):
                pictures += 1
            if sh.has_text_frame and sh.text_frame.text.strip():
                boxes += 1
                chars += len(sh.text_frame.text)
    assert boxes > 50, f"可编辑文本框太少（{boxes}），可能退化成了图片"
    assert chars > 1000, f"可编辑字符太少（{chars}）"
    assert pictures == 0, "这份 fixture 不含位图，出现 PICTURE 说明文字被光栅化了"


@needs_soffice
def test_known_content_survived(converted):
    """抽查几段原文，确认内容没丢。"""
    all_text = "\n".join(
        sh.text_frame.text
        for slide in converted.slides
        for sh in slide.shapes
        if sh.has_text_frame
    )
    for needle in ["示例科技", "全年排期", "擅长领域", "Sample Corp", "09:30-09:50"]:
        assert needle in all_text, f"内容丢失：{needle}"


@needs_soffice
def test_fix1_font_names_are_powerpoint_resolvable(converted, raw_converted):
    """缺陷 1：LibreOffice 写无空格字体名，PowerPoint 认不出。"""
    def names(prs):
        return {
            r.font.name
            for s in prs.slides for sh in s.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs
            if r.font.name
        }

    raw, fixed = names(raw_converted), names(converted)
    assert raw != fixed, "修补前后字体名完全一致 —— fixture 没有触发这条修补"
    for n in fixed:
        assert n not in C.FONT_RENAME, f"仍残留无空格字体名：{n}"
        assert not C.SUBSET_PREFIX.match(n), f"仍残留子集前缀：{n}"


@needs_soffice
def test_fix1b_cjk_runs_get_east_asian_typeface(converted):
    """缺陷 1b：只写 a:latin 不写 a:ea，中文没有兜底字体。"""
    missing = 0
    checked = 0
    for slide in converted.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text and C.CJK.search(r.text):
                        checked += 1
                        rPr = r._r.find(qn("a:rPr"))
                        ea = rPr.find(qn("a:ea")) if rPr is not None else None
                        if ea is None or not ea.get("typeface"):
                            missing += 1
    assert checked > 20, "fixture 里中文 run 太少，测不出东亚字体这条"
    assert missing == 0, f"{missing}/{checked} 个中文 run 没有 a:ea 东亚字体"


@needs_soffice
def test_fix2_wrap_disabled_everywhere(converted, raw_converted):
    """缺陷 2：文本框没关自动换行，字体一替换就折行叠字。"""
    def wrap_states(prs):
        return [
            sh.text_frame.word_wrap
            for s in prs.slides for sh in s.shapes if sh.has_text_frame
        ]

    assert any(w is not False for w in wrap_states(raw_converted)), \
        "修补前就已经全部关掉了 —— fixture 没有触发这条修补"
    assert all(w is False for w in wrap_states(converted)), "仍有文本框开着自动换行"


@needs_soffice
def test_cjk_font_override(tmp_path):
    out = tmp_path / "deck.pptx"
    subprocess.run(
        [sys.executable, str(CONVERT), str(FIXTURE), "-o", str(out),
         "--cjk-font", "Noto Sans SC"],
        check=True, capture_output=True, timeout=600,
    )
    faces = set()
    for s in Presentation(str(out)).slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text and C.CJK.search(r.text):
                        rPr = r._r.find(qn("a:rPr"))
                        ea = rPr.find(qn("a:ea")) if rPr is not None else None
                        if ea is not None:
                            faces.add(ea.get("typeface"))
    assert faces == {"Noto Sans SC"}, f"--cjk-font 没生效：{faces}"


# ------------------------------------------------------------- CLI 行为
def test_rejects_non_pdf(tmp_path):
    f = tmp_path / "not.txt"
    f.write_text("hi")
    r = subprocess.run([sys.executable, str(CONVERT), str(f)], capture_output=True, text=True)
    assert r.returncode != 0 and "不是 PDF" in r.stderr


def test_rejects_missing_file():
    r = subprocess.run([sys.executable, str(CONVERT), "/nope.pdf"], capture_output=True, text=True)
    assert r.returncode != 0 and "找不到文件" in r.stderr


def test_rejects_bad_font_map():
    r = subprocess.run(
        [sys.executable, str(CONVERT), str(FIXTURE), "--font-map", "oops"],
        capture_output=True, text=True,
    )
    assert r.returncode != 0 and "--font-map" in r.stderr


def test_doctor_runs():
    r = subprocess.run([sys.executable, str(CONVERT), "--doctor"], capture_output=True, text=True)
    assert "LibreOffice" in r.stdout and "python-pptx" in r.stdout


# -------------------------------------------- 缺陷 3/4：文字内容订正
def test_radical_table_targets_are_real_ideographs():
    """对照表打错字比不修更糟 —— 校验每个目标都是真正的 CJK 统一表意文字。"""
    import unicodedata as ud

    for radical, target in C.RADICAL_SUPPLEMENT.items():
        assert len(radical) == 1 and len(target) == 1, f"{radical!r}->{target!r} 不是单字符"
        assert 0x2E80 <= ord(radical) <= 0x2EFF, f"{radical!r} 不在部首补充区"
        assert 0x4E00 <= ord(target) <= 0x9FFF, f"{target!r} 不是 CJK 统一表意文字"
        assert ud.name(target).startswith("CJK UNIFIED IDEOGRAPH"), f"{target!r} 名字不对"
    assert len(set(C.RADICAL_SUPPLEMENT.values())) == len(C.RADICAL_SUPPLEMENT), "目标有重复"


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("⽰例科技", "示例科技"),            # 康熙部首，NFKC 可解
        ("擅⻓领域", "擅长领域"),            # 部首补充区，靠对照表
        ("⻛险清单", "风险清单"),
        ("正常文字", "正常文字"),                 # 不该动
        ("ABC 123", "ABC 123"),
        ("全角：（）", "全角：（）"),              # 全角标点不许被 NFKC 拉成半角
    ],
)
def test_unkangxi(raw, expected):
    assert C.unkangxi(raw)[0] == expected


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("擅 长 领 域：", "擅长领域："),          # 字距 artefact
        ("三 种 模 式", "三种模式"),
        ("这是 正常 中文", "这是 正常 中文"),      # 词间空格，不该合并
        ("a b c", "a b c"),                      # 拉丁文，不该合并
        ("中 文", "中 文"),                       # 只有两字，证据不足，不动
        ("", ""),
    ],
)
def test_unspace_cjk(raw, expected):
    assert C.unspace_cjk(raw)[0] == expected


@needs_soffice
def test_fix3_no_compat_chars_left(converted):
    """缺陷 3：兼容字符看着一样但搜不到，必须还原成常规汉字。"""
    leftovers = []
    for slide in converted.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                leftovers += [c for c in sh.text_frame.text if 0x2E80 <= ord(c) <= 0x2FDF]
    assert not leftovers, f"仍残留兼容字符：{set(leftovers)}"


@needs_soffice
def test_fix3_searchable_after_cleanup(converted, raw_converted):
    """订正前搜不到、订正后搜得到 —— 证明这条修补确实有价值。"""
    def text(prs):
        return "\n".join(
            sh.text_frame.text
            for s in prs.slides for sh in s.shapes if sh.has_text_frame
        )

    before, after = text(raw_converted), text(converted)
    for needle in ["风险清单", "擅长领域", "目标对齐", "骨架固定"]:
        assert needle not in before, f"fixture 没触发这条修补：{needle} 修补前就能搜到"
        assert needle in after, f"修补后仍搜不到：{needle}"
