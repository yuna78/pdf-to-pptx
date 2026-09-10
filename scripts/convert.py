#!/usr/bin/env python3
"""PDF → 可编辑 PPTX。

链路：LibreOffice 的 PDF 导入过滤器把 PDF 还原成 Impress 文档（文字变真文本框、
矢量路径变 FREEFORM 形状），导出 pptx，再用 python-pptx 修掉 LibreOffice 的两个
系统性缺陷（字体名 / 自动换行），最后可选渲染回 PDF 做逐页比对。

用法见同目录 ../SKILL.md。
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from collections import Counter
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
VENV_DIR = SKILL_DIR / ".venv"
REQUIRED = ["python-pptx", "pillow"]


# ---------------------------------------------------------------- bootstrap
def ensure_deps() -> None:
    """确保跑在一个装了 python-pptx + pillow 的解释器里，否则建 venv 重入。"""
    try:
        import pptx  # noqa: F401
        import PIL  # noqa: F401
        return
    except ImportError:
        pass

    venv_py = VENV_DIR / "bin" / "python"
    if not venv_py.exists():
        print("[setup] 首次运行，创建 venv 并安装 python-pptx / pillow …", file=sys.stderr)
        subprocess.run([sys.executable, "-m", "venv", str(VENV_DIR)], check=True)
        subprocess.run(
            [str(venv_py), "-m", "pip", "install", "-q", "--upgrade", "pip", *REQUIRED],
            check=True,
        )
    if Path(sys.executable).resolve() == venv_py.resolve():
        print("venv 里仍然缺依赖，请删掉 %s 重跑" % VENV_DIR, file=sys.stderr)
        sys.exit(1)
    os.execv(str(venv_py), [str(venv_py), os.path.abspath(__file__), *sys.argv[1:]])


# ------------------------------------------------------------------- tools
SOFFICE_CANDIDATES = [
    "soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/lib/libreoffice/program/soffice",
]


def find_soffice(explicit: str | None = None) -> str:
    if explicit:
        if Path(explicit).exists() or shutil.which(explicit):
            return explicit
        die(f"指定的 soffice 不存在：{explicit}")
    for c in SOFFICE_CANDIDATES:
        p = shutil.which(c) if not c.startswith("/") else (c if Path(c).exists() else None)
        if p:
            return p
    die("找不到 LibreOffice。macOS: brew install --cask libreoffice / Linux: apt install libreoffice-impress")


def die(msg: str) -> None:
    print(f"错误：{msg}", file=sys.stderr)
    sys.exit(1)


def doctor() -> int:
    ok = True
    rows = []
    try:
        s = find_soffice()
        rows.append(("LibreOffice", "OK", s))
    except SystemExit:
        rows.append(("LibreOffice", "缺失", "brew install --cask libreoffice"))
        ok = False
    for tool, hint in (("pdftoppm", "brew install poppler"), ("pdfinfo", "brew install poppler")):
        p = shutil.which(tool)
        rows.append((tool, "OK" if p else "缺失（只影响 --verify）", p or hint))
    try:
        import pptx  # noqa: F401

        rows.append(("python-pptx", "OK", ""))
    except ImportError:
        rows.append(("python-pptx", "缺失", "首次运行会自动建 venv"))
    for name, status, extra in rows:
        print(f"{name:<14} {status:<22} {extra}")
    return 0 if ok else 1


# ------------------------------------------------------------- step 1: 转换
def pdf_to_pptx_raw(pdf: Path, outdir: Path, soffice: str, timeout: int) -> Path:
    """LibreOffice PDF 导入 → pptx。用独立 profile，避免撞上已开着的 LibreOffice。"""
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir / "lo-profile"
    cmd = [
        soffice,
        "--headless",
        "--norestore",
        f"-env:UserInstallation=file://{profile}",
        "--infilter=impress_pdf_import",
        "--convert-to",
        "pptx",
        "--outdir",
        str(outdir),
        str(pdf),
    ]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    produced = outdir / (pdf.stem + ".pptx")
    if not produced.exists():
        die(
            "LibreOffice 没有产出 pptx。\n"
            f"stdout: {r.stdout.strip()}\nstderr: {r.stderr.strip()}\n"
            "常见原因：LibreOffice 正开着（本脚本已用独立 profile，仍失败就先退出 LibreOffice）、"
            "或该 PDF 加密受保护。"
        )
    return produced


# ------------------------------------------------------------- step 2: 修补
# LibreOffice 的 PDF 导入把字体名写成无空格形式，PowerPoint 认不出。
# 只收录「真名带空格」的；SimSun / SimHei / KaiTi / FangSong 等本来就没空格，不要动。
FONT_RENAME = {
    "MicrosoftYaHei": "Microsoft YaHei",
    "MicrosoftYaHeiUI": "Microsoft YaHei UI",
    "MicrosoftJhengHei": "Microsoft JhengHei",
    "MicrosoftJhengHeiUI": "Microsoft JhengHei UI",
    "MicrosoftSansSerif": "Microsoft Sans Serif",
    "TimesNewRoman": "Times New Roman",
    "CourierNew": "Courier New",
    "ArialUnicodeMS": "Arial Unicode MS",
    "SegoeUI": "Segoe UI",
    "SourceHanSansCN": "Source Han Sans CN",
    "SourceHanSerifCN": "Source Han Serif CN",
    "NotoSansSC": "Noto Sans SC",
    "NotoSerifSC": "Noto Serif SC",
    "PingFangSC": "PingFang SC",
    "HiraginoSansGB": "Hiragino Sans GB",
}
SUBSET_PREFIX = re.compile(r"^[A-Z]{6}\+")
CJK = re.compile(r"[㐀-鿿豈-﫿぀-ヿ가-힯]")


# 缺陷 3：部分汉字被还原成「康熙部首」等兼容字符（⽰ U+2F70 而不是 示 U+793A）。
# 看着一模一样，但复制出去搜不到。只处理这两段，不动全角标点。
COMPAT_RANGES = re.compile("[\u2e80-\u2fdf\uf900-\ufaff]")
# 记下还原不了的兼容字符，跑完统一警告 —— 宁可报出来，也不瞎猜。
UNMAPPED_COMPAT: "Counter[str]" = Counter()
# 缺陷 4：源 PDF 用了字距调整时，LibreOffice 会在每个汉字之间塞一个空格
# （「擅长领域」变成「擅 长 领 域」）。只在「每个汉字都被单空格隔开」这种
# 明显是字距 artefact 的形态下合并；正常中文不会这么写。
_C = "[\u3400-\u9fff\uf900-\ufaff]"
SPACED_CJK = re.compile("^(?:%s ){2,}%s(?P<tail>[\uff1a:\uff0c\u3002\u3001\uff01\uff1f\uff1b\uff09\u3011\u300d\u300b]*)$" % (_C, _C))


# CJK 部首补充区（U+2E80–U+2EFF）没有 NFKC 分解，只能靠对照表。
# 这里只收「整字型」部首（本身就是一个常用汉字的简化形），逐条人工核对过。
# 表外的字符不猜 —— 脚本会警告并原样保留，欢迎提 PR 补充。
RADICAL_SUPPLEMENT = {
    "\u2ec4": "\u897f",  # ⻄ WEST -> 西
    "\u2ec5": "\u89c1",  # ⻅ C-SIMPLIFIED SEE -> 见
    "\u2ec9": "\u8d1d",  # ⻉ C-SIMPLIFIED SHELL -> 贝
    "\u2ecb": "\u8f66",  # ⻋ C-SIMPLIFIED CART -> 车
    "\u2ed3": "\u957f",  # ⻓ C-SIMPLIFIED LONG -> 长
    "\u2ed4": "\u95e8",  # ⻔ C-SIMPLIFIED GATE -> 门
    "\u2ed8": "\u9752",  # ⻘ BLUE -> 青
    "\u2ed9": "\u97e6",  # ⻙ C-SIMPLIFIED TANNED LEATHER -> 韦
    "\u2eda": "\u9875",  # ⻚ C-SIMPLIFIED LEAF -> 页
    "\u2edb": "\u98ce",  # ⻛ C-SIMPLIFIED WIND -> 风
    "\u2edc": "\u98de",  # ⻜ C-SIMPLIFIED FLY -> 飞
    "\u2edd": "\u98df",  # ⻝ EAT -> 食
    "\u2ee2": "\u9a6c",  # ⻢ C-SIMPLIFIED HORSE -> 马
    "\u2ee3": "\u9aa8",  # ⻣ BONE -> 骨
    "\u2ee4": "\u9b3c",  # ⻤ GHOST -> 鬼
    "\u2ee5": "\u9c7c",  # ⻥ C-SIMPLIFIED FISH -> 鱼
    "\u2ee6": "\u9e1f",  # ⻦ C-SIMPLIFIED BIRD -> 鸟
    "\u2eec": "\u9f50",  # ⻬ C-SIMPLIFIED EVEN -> 齐
    "\u2eee": "\u9f7f",  # ⻮ C-SIMPLIFIED TOOTH -> 齿
    "\u2ef0": "\u9f99",  # ⻰ C-SIMPLIFIED DRAGON -> 龙
    "\u2ef3": "\u9f9f",  # ⻳ C-SIMPLIFIED TURTLE -> 龟
}


def unkangxi(text):
    """康熙部首 / 兼容汉字 → 常规汉字。只在 NFKC 结果是单个字符时才换。"""
    if not COMPAT_RANGES.search(text):
        return text, 0
    out, n = [], 0
    for ch in text:
        if COMPAT_RANGES.match(ch):
            norm = RADICAL_SUPPLEMENT.get(ch) or unicodedata.normalize("NFKC", ch)
            if len(norm) == 1 and norm != ch:
                out.append(norm)
                n += 1
                continue
            UNMAPPED_COMPAT[ch] += 1
        out.append(ch)
    return "".join(out), n


def unspace_cjk(text):
    """「擅 长 领 域」→「擅长领域」。形态不匹配就原样返回。"""
    stripped = text.strip()
    m = SPACED_CJK.match(stripped)
    if not m:
        return text, 0
    tail = m.group("tail")
    body = stripped[: len(stripped) - len(tail)] if tail else stripped
    merged = body.replace(" ", "") + tail
    if merged == stripped:
        return text, 0
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    return lead + merged + trail, 1


def norm_font(name: str | None) -> str | None:
    if not name:
        return name
    name = SUBSET_PREFIX.sub("", name)
    return FONT_RENAME.get(name, name)


def iter_runs(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    yield run


def dominant_cjk_font(prs) -> str | None:
    """全篇承载中文最多的字体，用作 a:ea（东亚字体）。"""
    c: Counter = Counter()
    for run in iter_runs(prs):
        if run.text and CJK.search(run.text):
            n = norm_font(run.font.name)
            if n:
                c[n] += len(run.text)
    return c.most_common(1)[0][0] if c else None


def postfix(src: Path, dst: Path, *, keep_wrap: bool, cjk_font: str | None,
            extra_map: dict[str, str], text_cleanup: bool = True) -> dict:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation(str(src))
    FONT_RENAME.update(extra_map)

    ea_font = cjk_font or dominant_cjk_font(prs)
    stats = {"font_renamed": 0, "ea_set": 0, "wrap_off": 0, "autofit_off": 0,
             "kangxi": 0, "unspaced": 0}

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame

            # 缺陷 2：文本框宽度按原字体量的，字体一被替换就折行叠字 → 关掉自动换行。
            if not keep_wrap:
                if tf.word_wrap is not False:
                    tf.word_wrap = False
                    stats["wrap_off"] += 1
                bodyPr = tf._txBody.bodyPr
                for tag in ("a:normAutofit", "a:spAutoFit"):
                    for el in bodyPr.findall(qn(tag)):
                        bodyPr.remove(el)
                        stats["autofit_off"] += 1

            for para in tf.paragraphs:
                for run in para.runs:
                    # 缺陷 1a：字体名无空格，PowerPoint 认不出。
                    old = run.font.name
                    new = norm_font(old)
                    if new and new != old:
                        run.font.name = new
                        stats["font_renamed"] += 1
                    # 缺陷 3 / 4：康熙部首兼容字符、字距造成的逐字空格。
                    #（改文字要在设 a:ea 之前做，否则 run.text 判断的是旧值）
                    if text_cleanup and run.text:
                        t, n1 = unkangxi(run.text)
                        t, n2 = unspace_cjk(t)
                        if n1 or n2:
                            run.text = t
                            stats["kangxi"] += n1
                            stats["unspaced"] += n2
                    # 缺陷 1b：只写了 a:latin，中文那一半没有 a:ea 兜底。
                    if ea_font and run.text and CJK.search(run.text):
                        rPr = run._r.get_or_add_rPr()
                        ea = rPr.find(qn("a:ea"))
                        if ea is None:
                            ea = rPr.makeelement(qn("a:ea"), {})
                            rPr.append(ea)
                        if ea.get("typeface") != ea_font:
                            ea.set("typeface", ea_font)
                            stats["ea_set"] += 1

    prs.save(str(dst))
    stats["cjk_font"] = ea_font or "(未检出中文)"
    return stats


# ------------------------------------------------------------- step 3: 体检
def inspect(pptx_path: Path) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        row = {"page": i, "shapes": 0, "textboxes": 0, "chars": 0, "pictures": 0, "vectors": 0}
        for sh in slide.shapes:
            row["shapes"] += 1
            st = str(sh.shape_type or "")
            if "PICTURE" in st:
                row["pictures"] += 1
            elif "FREEFORM" in st or "AUTO_SHAPE" in st:
                row["vectors"] += 1
            if sh.has_text_frame and sh.text_frame.text.strip():
                row["textboxes"] += 1
                row["chars"] += len(sh.text_frame.text)
        out.append(row)
    return out


def has_text_layer(pdf: Path) -> bool:
    """扫描件（无文字层）转出来只会是一堆图片，提前警告。"""
    exe = shutil.which("pdftotext")
    if not exe:
        return True  # 无法判断，不误报
    r = subprocess.run([exe, "-l", "3", str(pdf), "-"], capture_output=True, text=True)
    return len(r.stdout.strip()) > 30


# ------------------------------------------------------------- step 4: 比对
def verify(pdf: Path, pptx_path: Path, outdir: Path, soffice: str, dpi: int,
           pages: list[int] | None, timeout: int) -> list[Path]:
    from PIL import Image, ImageDraw

    if not shutil.which("pdftoppm"):
        print("[verify] 跳过：没装 pdftoppm（brew install poppler）", file=sys.stderr)
        return []
    outdir.mkdir(parents=True, exist_ok=True)
    back = outdir / "back"
    back.mkdir(exist_ok=True)
    profile = outdir / "lo-profile-verify"
    subprocess.run(
        [soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{profile}",
         "--convert-to", "pdf", "--outdir", str(back), str(pptx_path)],
        capture_output=True, text=True, timeout=timeout,
    )
    back_pdf = back / (pptx_path.stem + ".pdf")
    if not back_pdf.exists():
        print("[verify] 跳过：pptx 渲染回 PDF 失败", file=sys.stderr)
        return []

    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), str(pdf), str(outdir / "orig")], check=True)
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), str(back_pdf), str(outdir / "conv")], check=True)

    origs = sorted(outdir.glob("orig-*.png"))
    convs = sorted(outdir.glob("conv-*.png"))
    made = []
    for idx, (a_p, b_p) in enumerate(zip(origs, convs), 1):
        if pages and idx not in pages:
            continue
        a, b = Image.open(a_p).convert("RGB"), Image.open(b_p).convert("RGB")
        c = Image.new("RGB", (a.width + b.width + 16, max(a.height, b.height) + 24), "white")
        c.paste(a, (0, 24))
        c.paste(b, (a.width + 16, 24))
        d = ImageDraw.Draw(c)
        d.text((6, 6), f"p{idx} ORIGINAL PDF", fill="black")
        d.text((a.width + 22, 6), f"p{idx} CONVERTED PPTX", fill="red")
        out = outdir / f"cmp-{idx:02d}.png"
        c.save(out)
        made.append(out)
    if len(origs) != len(convs):
        print(f"[verify] 注意：原稿 {len(origs)} 页，产物 {len(convs)} 页，页数不一致", file=sys.stderr)
    return made


# -------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser(description="PDF → 可编辑 PPTX")
    ap.add_argument("pdf", nargs="?", help="输入 PDF")
    ap.add_argument("-o", "--out", help="输出 .pptx 路径或目录；默认与 PDF 同目录同名")
    ap.add_argument("--verify", action="store_true", help="渲染回 PDF 并生成逐页左右对照图")
    ap.add_argument("--verify-dir", help="对照图输出目录；默认临时目录")
    ap.add_argument("--verify-pages", help="只比对指定页，如 1,5,6")
    ap.add_argument("--verify-dpi", type=int, default=80, help="对照图分辨率，默认 80")
    ap.add_argument("--cjk-font", help="强制指定东亚字体（a:ea），如 'PingFang SC'")
    ap.add_argument("--font-map", action="append", default=[],
                    help="额外字体改名，可重复：--font-map 'FooBar=Foo Bar'")
    ap.add_argument("--keep-wrap", action="store_true",
                    help="保留 LibreOffice 的自动换行设置（默认关掉，防叠字）")
    ap.add_argument("--no-text-cleanup", action="store_true",
                    help="不修文字内容（保留康熙部首兼容字符与逐字空格）")
    ap.add_argument("--no-postfix", action="store_true", help="不做任何修补，交出 LibreOffice 原始产物")
    ap.add_argument("--soffice", help="指定 soffice 路径")
    ap.add_argument("--timeout", type=int, default=600, help="单次 LibreOffice 调用超时秒数")
    ap.add_argument("--doctor", action="store_true", help="检查依赖")
    args = ap.parse_args()

    if args.doctor:
        return doctor()
    if not args.pdf:
        ap.error("缺少输入 PDF（或用 --doctor 检查依赖）")

    pdf = Path(args.pdf).expanduser().resolve()
    if not pdf.exists():
        die(f"找不到文件：{pdf}")
    if pdf.suffix.lower() != ".pdf":
        die(f"输入不是 PDF：{pdf.name}")

    if args.out:
        out = Path(args.out).expanduser().resolve()
        if out.is_dir() or not out.suffix:
            out = out / (pdf.stem + ".pptx")
    else:
        out = pdf.with_suffix(".pptx")
    out.parent.mkdir(parents=True, exist_ok=True)

    soffice = find_soffice(args.soffice)

    if not has_text_layer(pdf):
        print("[警告] 这份 PDF 前几页几乎取不到文字，很可能是扫描件。"
              "转出来只会是图片，改不了字——需要先 OCR。", file=sys.stderr)

    extra_map = {}
    for m in args.font_map:
        if "=" not in m:
            die(f"--font-map 格式应为 '原名=新名'，收到：{m}")
        k, v = m.split("=", 1)
        extra_map[k.strip()] = v.strip()

    with tempfile.TemporaryDirectory(prefix="pdf2pptx-") as tmp:
        tmpdir = Path(tmp)
        raw = pdf_to_pptx_raw(pdf, tmpdir / "raw", soffice, args.timeout)

        if args.no_postfix:
            shutil.copy2(raw, out)
            stats = None
        else:
            stats = postfix(raw, out, keep_wrap=args.keep_wrap,
                            cjk_font=args.cjk_font, extra_map=extra_map,
                            text_cleanup=not args.no_text_cleanup)

        rows = inspect(out)
        print(f"\n页数 {len(rows)}")
        print(f"{'页':>3}  {'形状':>5} {'文本框':>6} {'字符':>6} {'矢量':>5} {'图片':>5}")
        for r in rows:
            print(f"{r['page']:>3}  {r['shapes']:>5} {r['textboxes']:>6} {r['chars']:>6} "
                  f"{r['vectors']:>5} {r['pictures']:>5}")
        total_chars = sum(r["chars"] for r in rows)
        if total_chars == 0:
            print("\n[警告] 产物里一个可编辑文字都没有——大概率是扫描件或纯图 PDF。", file=sys.stderr)
        if stats:
            print(f"\n修补：字体改名 {stats['font_renamed']} 处 · 东亚字体 {stats['ea_set']} 处 "
                  f"· 关自动换行 {stats['wrap_off']} 个框 · 去 autofit {stats['autofit_off']} 处")
            if stats["kangxi"] or stats["unspaced"]:
                print(f"文字订正：部首/兼容字符还原 {stats['kangxi']} 个 "
                      f"· 逐字空格合并 {stats['unspaced']} 处")
            if UNMAPPED_COMPAT:
                items = " ".join(f"{c}(U+{ord(c):04X})×{n}"
                                 for c, n in UNMAPPED_COMPAT.most_common(12))
                print(f"\n[警告] 有 {sum(UNMAPPED_COMPAT.values())} 个兼容字符没有对照表、"
                      f"无法自动还原，已原样保留：\n  {items}\n"
                      "  它们看着和正常汉字一样，但复制出去搜不到。"
                      "去 scripts/convert.py 的 RADICAL_SUPPLEMENT 补一条即可。",
                      file=sys.stderr)
            print(f"东亚字体设为：{stats['cjk_font']}")

        if args.verify:
            vdir = Path(args.verify_dir).expanduser().resolve() if args.verify_dir \
                else tmpdir / "verify"
            pages = [int(x) for x in args.verify_pages.split(",")] if args.verify_pages else None
            imgs = verify(pdf, out, vdir, soffice, args.verify_dpi, pages, args.timeout)
            if imgs:
                print("\n对照图（用 Read 逐张看，别只看上面的数字）：")
                for p in imgs:
                    print(f"  {p}")
                if not args.verify_dir:
                    keep = out.parent / f"{out.stem}-verify"
                    shutil.rmtree(keep, ignore_errors=True)
                    shutil.copytree(vdir, keep)
                    print(f"\n对照图已复制到：{keep}")

    print(f"\nPPTX: {out}")
    return 0


if __name__ == "__main__":
    ensure_deps()
    sys.exit(main())
